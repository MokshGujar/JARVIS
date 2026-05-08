from __future__ import annotations

import logging
import csv
import os
import re
import time
import urllib.parse
from pathlib import Path
from typing import Callable, Dict, Iterable

from config import BASE_DIR as CONFIG_BASE_DIR
from app.orchestrator.action_plan import ActionPlan, ActionStep
from app.orchestrator.tool_executor import ToolExecutor
from app.tools.automation_domain_helper import ServiceBackedDomainHelper
from app.services.contact_match_service import ContactCandidate, ContactMatchService
from app.services.message_action_service import MessageActionService
from app.tools.base import ToolContext

try:
    from send2trash import send2trash
    SEND2TRASH_IMPORT_ERROR = None
except Exception as exc:
    send2trash = None
    SEND2TRASH_IMPORT_ERROR = exc


logger = logging.getLogger("J.A.R.V.I.S")


SEARCH_EXCLUDED_DIR_NAMES = {
    "$recycle.bin",
    ".cache",
    ".continue",
    ".git",
    ".gradle",
    ".idea",
    ".mypy_cache",
    ".pytest_cache",
    ".ruff_cache",
    ".tox",
    ".venv",
    "__pycache__",
    "appdata",
    "cache",
    "caches",
    "database",
    "dist",
    "env",
    "logs",
    "model_cache",
    "models",
    "node_modules",
    "pytest-cache-files",
    "site-packages",
    "temp",
    "tmp",
    "venv",
}

SEARCH_TEXT_EXTENSIONS = {
    ".bat",
    ".cfg",
    ".csv",
    ".css",
    ".html",
    ".ini",
    ".js",
    ".json",
    ".log",
    ".md",
    ".py",
    ".ps1",
    ".toml",
    ".ts",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}

READ_TEXT_EXTENSIONS = SEARCH_TEXT_EXTENSIONS | {
    ".c",
    ".cpp",
    ".cs",
    ".go",
    ".java",
    ".kt",
    ".php",
    ".rs",
    ".sh",
    ".sql",
}

MAX_FILE_SEARCH_SECONDS = 4.0
MAX_FILE_SEARCH_SCANNED = 12_000
MAX_CONTENT_SEARCH_BYTES = 256_000


def _runtime_base_dir() -> Path:
    try:
        from app.services import automation_service as automation_module

        return automation_module.BASE_DIR
    except Exception:
        return CONFIG_BASE_DIR



class AutomationFileCompatibility(ServiceBackedDomainHelper):

    def _sanitize_file_reference(self, path_text: str) -> str:
            cleaned = (path_text or "").strip().strip('"').strip("'")
            cleaned = re.sub(r"\bin it\b", "", cleaned, flags=re.IGNORECASE)
            cleaned = re.sub(r"\b(?:uh|um|er|ah)\b", "", cleaned, flags=re.IGNORECASE)
            cleaned = re.sub(r"^(?:for me\s+)?(?:please\s+)?", "", cleaned, flags=re.IGNORECASE)
            cleaned = re.sub(r"\bdot\s+([a-z0-9]{1,5})\b", r".\1", cleaned, flags=re.IGNORECASE)
            cleaned = re.sub(r"\s+", " ", cleaned).strip()
            return cleaned.rstrip(".!?")


    def _strip_named_prefix(self, value: str) -> str:
            value = re.sub(r"^(?:named|called|name)\s+(?:the\s+)?", "", value, flags=re.IGNORECASE)
            return value.strip()


    def _clean_location_phrase(self, value: str) -> str:
            cleaned = (value or "").strip().strip('"').strip("'")
            cleaned = re.sub(
                r"^(?:the\s+)?(?:folder\s+)?(?:path\s+)?(?:location\s+)?",
                "",
                cleaned,
                flags=re.IGNORECASE,
            )
            cleaned = re.sub(
                r"^(?:on|in|at|inside|under|from)\s+",
                "",
                cleaned,
                flags=re.IGNORECASE,
            )
            cleaned = re.sub(r"\s+", " ", cleaned).strip()
            return cleaned.rstrip(".!?")


    def _clean_file_name(self, value: str) -> str:
            cleaned = self._strip_named_prefix(value or "")
            cleaned = re.sub(r"^the\s+", "", cleaned, flags=re.IGNORECASE)
            cleaned = re.sub(r"^(?:the\s+)?(?:file|folder|directory)\s+", "", cleaned, flags=re.IGNORECASE)
            cleaned = re.sub(r"\s+\.", ".", cleaned)
            cleaned = re.sub(r"\s+", " ", cleaned).strip()
            return cleaned.rstrip(".!?")


    def _looks_like_windows_absolute_path(self, value: str) -> bool:
            return bool(re.match(r'^[A-Za-z]:[\\/]', (value or "").strip()))


    def _combine_location_and_name(self, location_text: str, name_text: str) -> Path:
            location_path = self._resolve_path(self._clean_location_phrase(location_text))
            file_name = self._clean_file_name(name_text)
            return location_path / Path(file_name.replace("/", "\\"))


    def _extract_path_from_sentence(self, text: str) -> Path | None:
            cleaned = self._sanitize_file_reference(text)
            if not cleaned:
                return None

            absolute_match = re.search(r'([A-Za-z]:[\\/][^"\']*)$', cleaned)
            if absolute_match:
                return Path(absolute_match.group(1).strip().replace("/", "\\"))

            patterns = [
                r"^(?P<name>.+?)\s+(?:on|in|at|inside|under|from)\s+(?P<location>.+)$",
                r"^(?:on|in|at|inside|under|from)\s+(?P<location>.+?)\s+(?:named|called|name)\s+(?P<name>.+)$",
                r"^(?P<location>desktop|documents|downloads|pictures|videos|music|home)\s+(?:named|called|name)\s+(?P<name>.+)$",
                r"^(?:named|called|name)\s+(?P<name>.+?)\s+(?:on|in|at|inside|under|from)\s+(?P<location>.+)$",
            ]
            for pattern in patterns:
                match = re.match(pattern, cleaned, flags=re.IGNORECASE)
                if not match:
                    continue

                location = match.groupdict().get("location") or ""
                name = match.groupdict().get("name") or ""
                if location and name:
                    return self._combine_location_and_name(location, name)

            return None


    def _resolve_user_alias_path(self, path_text: str) -> Path | None:
            cleaned = self._clean_location_phrase(path_text)
            if not cleaned:
                return None

            for alias, base_path in self.USER_PATH_ALIASES.items():
                suffix_pattern = (
                    rf"^(?P<name>.+?)\s+(?:on|in|at|inside|under)\s+(?:the\s+)?{alias}(?:\s+folder)?$"
                )
                suffix_match = re.match(suffix_pattern, cleaned, flags=re.IGNORECASE)
                if suffix_match:
                    file_name = self._strip_named_prefix(suffix_match.group("name") or "")
                    if file_name:
                        return base_path / Path(file_name.replace("/", "\\"))

                pattern = (
                    rf"^(?:(?:on|in|at|inside|under)\s+(?:the\s+)?)?{alias}"
                    rf"(?:[\\/]|(?:\s+(?:folder\s+)?))?(.*)$"
                )
                match = re.match(pattern, cleaned, flags=re.IGNORECASE)
                if not match:
                    continue

                remainder = self._strip_named_prefix(match.group(1) or "")
                if not remainder:
                    return base_path

                remainder = remainder.lstrip("\\/ ").replace("/", "\\")
                return base_path / Path(remainder)

            return None


    def _resolve_path(self, path_text: str) -> Path:
            cleaned = self._sanitize_file_reference(path_text)
            extracted_path = self._extract_path_from_sentence(cleaned)
            if extracted_path is not None:
                return extracted_path

            alias_path = self._resolve_user_alias_path(cleaned)
            if alias_path is not None:
                return alias_path

            path_candidate = self._clean_location_phrase(cleaned)
            path = Path(path_candidate.replace("/", "\\"))
            if path.is_absolute() or self._looks_like_windows_absolute_path(path_candidate):
                return path

            return _runtime_base_dir() / path


    def _resolve_file_target(self, path_text: str) -> Path:
            cleaned = self._sanitize_file_reference(path_text)
            lowered = cleaned.lower()

            if lowered in {"that file", "the file", "it"}:
                if self._last_file_target is None:
                    raise ValueError("I don't know which file you mean yet. Tell me the file name once first.")
                return self._last_file_target

            return self._resolve_laptop_path(cleaned)


    def _resolve_folder_target(self, path_text: str) -> Path:
            cleaned = self._sanitize_file_reference(path_text)
            lowered = cleaned.lower()

            if lowered in {"that folder", "the folder", "that directory", "the directory", "it"}:
                if self._last_folder_target is None:
                    raise ValueError("I don't know which folder you mean yet. Tell me the folder name once first.")
                return self._last_folder_target

            return self._resolve_laptop_path(cleaned)


    def _resolve_existing_target(self, path_text: str, target_kind: str = "") -> Path:
            cleaned = self._sanitize_file_reference(path_text)
            lowered = cleaned.lower()

            if target_kind == "file" and lowered in {"that file", "the file", "it"}:
                if self._last_file_target is None:
                    raise ValueError("I don't know which file you mean yet. Tell me the file name once first.")
                return self._last_file_target

            if target_kind in {"folder", "directory"} and lowered in {"that folder", "the folder", "that directory", "the directory", "it"}:
                if self._last_folder_target is None:
                    raise ValueError("I don't know which folder you mean yet. Tell me the folder name once first.")
                return self._last_folder_target

            if lowered in {"it", "that", "that item", "the item"}:
                if self._last_file_target is not None and self._last_file_target.exists():
                    return self._last_file_target
                if self._last_folder_target is not None and self._last_folder_target.exists():
                    return self._last_folder_target
                raise ValueError("I don't know which file or folder you mean yet. Tell me the name once first.")

            path = self._resolve_laptop_path(cleaned)
            if target_kind == "file" and path.exists() and path.is_dir():
                raise ValueError("That is a folder, not a file.")
            if target_kind in {"folder", "directory"} and path.exists() and not path.is_dir():
                raise ValueError("That is a file, not a folder.")
            return path


    def _resolve_laptop_path(self, path_text: str) -> Path:
            path = self._resolve_path(path_text).resolve()
            if self._is_protected_path(path):
                raise ValueError("That location is protected. Jarvis cannot access Windows or critical system folders.")
            return path


    def _create_file_in_folder(self, folder_text: str, name_text: str, content: str) -> Dict[str, str | bool]:
            try:
                folder = self._resolve_folder_target(folder_text)
            except ValueError as exc:
                return {"success": False, "action": "create_file", "message": str(exc)}

            file_name = self._clean_file_name(name_text)
            if not file_name:
                return {"success": False, "action": "create_file", "message": "Tell me the file name you want to use."}

            path = folder / Path(file_name.replace("/", "\\"))
            return self._create_file(str(path), content)


    def _create_file_or_ask_for_location(self, path_text: str, content: str) -> Dict[str, str | bool]:
            cleaned = self._sanitize_file_reference(path_text)
            if not cleaned:
                return {"success": False, "action": "create_file", "message": "Tell me the file name you want to use."}

            if not self._looks_like_explicit_path_request(cleaned):
                file_name = self._clean_file_name(cleaned)
                if not file_name:
                    return {"success": False, "action": "create_file", "message": "Tell me the file name you want to use."}
                self._pending_create_file = {
                    "name": file_name,
                    "content": content or "",
                }
                return {
                    "success": False,
                    "action": "create_file_location_needed",
                    "message": f"Where should I save {file_name}?",
                }

            return self._create_file(cleaned, content)


    def _handle_create_file_location_followup(self, command: str) -> Dict[str, str | bool]:
            pending = self._pending_create_file
            if not pending:
                return {
                    "success": False,
                    "action": "create_file",
                    "message": "I don't have a pending file to save.",
                }

            reply = self._normalize_spoken_command(command)
            lowered = reply.lower().rstrip(".!?")
            if lowered in {"cancel", "stop", "never mind", "no", "skip"}:
                self._pending_create_file = None
                return {
                    "success": False,
                    "action": "create_file",
                    "message": "File creation cancelled.",
                }

            try:
                folder = self._resolve_folder_target(reply)
            except ValueError as exc:
                return {"success": False, "action": "create_file", "message": str(exc)}

            self._pending_create_file = None
            file_name = str(pending.get("name", "")).strip()
            content = str(pending.get("content", ""))
            steps = [
                ActionStep("step1", "file", "file", "create_file", {"parent": str(folder), "filename": file_name}),
            ]
            if content:
                steps.append(
                    ActionStep(
                        "step2",
                        "file",
                        "file",
                        "write_file",
                        {"path": "{step1.path}", "content": content, "overwrite": False},
                        depends_on=["step1"],
                    )
                )
                steps.append(
                    ActionStep(
                        "step3",
                        "file",
                        "file",
                        "verify_exists",
                        {"path": "{step1.path}", "expected_content": content},
                        depends_on=["step1", "step2"],
                    )
                )
            target_path = folder / Path(file_name.replace("/", "\\"))
            plan_command = f"create file {target_path}"
            executor = ToolExecutor(registry=self._build_automation_tool_registry(), enforce_policy=True)
            return executor.execute(
                ActionPlan(
                    original_text=plan_command,
                    steps=steps,
                    is_multistep=bool(content),
                ),
                ToolContext(
                    command=plan_command,
                    intent="file",
                    session_id=self._active_session_id,
                    request_id=self._active_turn_id,
                    payload={"turn_id": self._active_turn_id} if self._active_turn_id else {},
                    security_state={"step_up_verified": self._active_step_up_verified},
                ),
            )


    def _resolve_openable_path(self, target: str) -> Path | None:
            cleaned = self._sanitize_file_reference(target)
            lowered = cleaned.lower()

            if lowered in {"that folder", "the folder", "that directory", "the directory"}:
                return self._last_folder_target if self._last_folder_target and self._last_folder_target.exists() else None

            if lowered in {"that file", "the file"}:
                return self._last_file_target if self._last_file_target and self._last_file_target.exists() else None

            if lowered == "it":
                if self._last_folder_target and self._last_folder_target.exists():
                    return self._last_folder_target
                if self._last_file_target and self._last_file_target.exists():
                    return self._last_file_target
                return None

            if not self._looks_like_explicit_path_request(cleaned):
                return None

            try:
                path = self._resolve_laptop_path(cleaned)
            except ValueError:
                return None

            return path if path.exists() else None


    def _looks_like_explicit_path_request(self, target: str) -> bool:
            cleaned = (target or "").strip().lower()
            if not cleaned:
                return False

            if cleaned.startswith(("on ", "in ", "at ", "inside ", "under ")):
                return True

            if self._looks_like_windows_absolute_path(cleaned):
                return True

            if any(token in cleaned for token in ("\\", "/", ":")):
                return True

            if any(alias in cleaned for alias in self.USER_PATH_ALIASES):
                return True

            if any(keyword in cleaned for keyword in ("folder", "directory", "file")):
                return True

            return False


    def _display_file_name(self, path: Path) -> str:
            return path.name or str(path)


    def _display_target_name(self, path: Path) -> str:
            return path.name or str(path)


    def _display_parent_name(self, path: Path) -> str:
            parent = path.parent
            try:
                resolved_parent = parent.resolve()
                for alias, alias_path in self.USER_PATH_ALIASES.items():
                    try:
                        if resolved_parent == alias_path.resolve():
                            return alias.capitalize()
                    except Exception:
                        continue
                base_dir = _runtime_base_dir()
                if resolved_parent == base_dir.resolve():
                    return base_dir.name
            except Exception:
                pass
            return parent.name or str(parent)


    def _remember_target(self, path: Path) -> None:
            if path.exists() and path.is_dir():
                self._last_folder_target = path
            else:
                self._last_file_target = path
                self._last_selected_file_path = path


    def _resolve_recent_file_selection(self, reference: str) -> Path | None:
            text = self._sanitize_file_reference(reference).lower()
            if text in {"it", "that", "this", "that file", "the file"}:
                selected = getattr(self, "_last_selected_file_path", None) or self._last_file_target
                return Path(selected) if selected else None
            index = self._ordinal_to_index(text)
            results = list(getattr(self, "_last_file_search_results", []) or [])
            if index == -1:
                index = len(results) - 1
            if index is None or index < 0 or index >= len(results):
                return None
            path_text = str(dict(results[index]).get("path") or "").strip()
            if not path_text:
                return None
            path = Path(path_text)
            self._last_selected_file_path = path
            self._last_file_target = path
            return path


    @staticmethod
    def _ordinal_to_index(text: str) -> int | None:
            normalized = re.sub(r"\b(?:the|one|file)\b", " ", str(text or "").lower())
            normalized = re.sub(r"\s+", " ", normalized).strip()
            mapping = {
                "first": 0,
                "1": 0,
                "second": 1,
                "2": 1,
                "third": 2,
                "3": 2,
                "fourth": 3,
                "4": 3,
                "fifth": 4,
                "5": 4,
                "last": -1,
            }
            return mapping.get(normalized)


    def _show_selected_file_path(self, reference: str = "it") -> Dict[str, object]:
            path = self._resolve_recent_file_selection(reference)
            if path is None:
                return {
                    "success": False,
                    "action": "show_file_path",
                    "status": "clarification_required",
                    "requires_followup": True,
                    "message": "Which file should I use?",
                }
            return {
                "success": True,
                "action": "show_file_path",
                "message": str(path),
                "path": str(path),
                "data": {"path": str(path)},
            }


    def _create_file(self, path_text: str, content: str) -> Dict[str, str | bool]:
            try:
                path = self._resolve_file_target(path_text)
            except ValueError as exc:
                return {"success": False, "action": "create_file", "message": str(exc)}

            if path.exists() and path.is_dir():
                return {
                    "success": False,
                    "action": "create_file",
                    "message": "That path is a folder, not a file.",
                }

            try:
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(content, encoding="utf-8")
            except Exception as exc:
                return {
                    "success": False,
                    "action": "create_file",
                    "message": f"Could not create {self._display_file_name(path)} at {path}: {exc}",
                }

            self._last_file_target = path
            return {
                "success": True,
                "action": "create_file",
                "message": f"Created {self._display_file_name(path)} in {self._display_parent_name(path)}.",
            }


    def _create_folder(self, path_text: str) -> Dict[str, str | bool]:
            try:
                path = self._resolve_folder_target(path_text)
            except ValueError as exc:
                return {"success": False, "action": "create_folder", "message": str(exc)}

            if path.exists() and not path.is_dir():
                return {
                    "success": False,
                    "action": "create_folder",
                    "message": "That path is a file, not a folder.",
                }

            try:
                path.mkdir(parents=True, exist_ok=True)
            except Exception as exc:
                return {
                    "success": False,
                    "action": "create_folder",
                    "message": f"Could not create folder {self._display_target_name(path)} at {path}: {exc}",
                }

            self._last_folder_target = path
            return {
                "success": True,
                "action": "create_folder",
                "message": f"Folder {self._display_target_name(path)} created at {path}.",
            }


    def create_file_with_content(self, path_text: str, content: str) -> Dict[str, str | bool]:
            return self._create_file(path_text, content)


    def _format_size(self, byte_count: int) -> str:
            size = float(max(byte_count, 0))
            for unit in ("B", "KB", "MB", "GB", "TB"):
                if size < 1024 or unit == "TB":
                    return f"{size:.1f} {unit}" if unit != "B" else f"{int(size)} B"
                size /= 1024
            return f"{size:.1f} TB"


    def _list_files(self, folder_text: str = "downloads", limit: int = 30) -> Dict[str, str | bool]:
            try:
                folder = self._resolve_folder_target(folder_text or "downloads")
            except ValueError as exc:
                return {"success": False, "action": "list_files", "message": str(exc)}

            if not folder.exists():
                return {"success": False, "action": "list_files", "message": f"{self._display_target_name(folder)} does not exist."}
            if not folder.is_dir():
                return {"success": False, "action": "list_files", "message": "That is a file, not a folder."}

            try:
                items = sorted(folder.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower()))
            except PermissionError:
                return {"success": False, "action": "list_files", "message": f"Permission denied: {folder}."}
            except Exception as exc:
                return {"success": False, "action": "list_files", "message": f"Could not list {folder}: {exc}"}

            visible_items = [item for item in items if not item.name.startswith(".")]
            lines = []
            for item in visible_items[:limit]:
                if item.is_dir():
                    lines.append(f"[folder] {item.name}/")
                else:
                    try:
                        size = self._format_size(item.stat().st_size)
                    except Exception:
                        size = "unknown size"
                    lines.append(f"[file] {item.name} ({size})")

            self._last_folder_target = folder
            if not lines:
                message = f"{self._display_target_name(folder)} is empty."
            else:
                suffix = f"\n... and {len(visible_items) - limit} more item(s)." if len(visible_items) > limit else ""
                message = f"Files in {self._display_target_name(folder)}:\n" + "\n".join(lines) + suffix
            return {"success": True, "action": "list_files", "message": message}


    def _read_file(self, path_text: str, max_chars: int = 4000) -> Dict[str, str | bool]:
            try:
                path = self._resolve_existing_target(path_text, target_kind="file")
            except ValueError as exc:
                return {"success": False, "action": "read_file", "message": str(exc)}

            if not path.exists():
                return {"success": False, "action": "read_file", "message": f"{self._display_file_name(path)} does not exist."}
            if not path.is_file():
                return {"success": False, "action": "read_file", "message": "That is a folder, not a file."}
            read_result = self._read_supported_file_content(path, max_chars=max_chars)
            if not read_result.get("success"):
                return read_result
            content = str(read_result.get("content") or "")

            self._last_file_target = path
            self._remember_target(path)
            return {
                "success": True,
                "action": "read_file",
                "message": f"{self._display_file_name(path)}:\n{content}" if content else f"{self._display_file_name(path)} is empty.",
                "path": str(path),
                "data": {"path": str(path), "content": content, "reader": read_result.get("reader")},
            }


    def _find_files(
        self,
        query_text: str,
        folder_text: str = "home",
        limit: int = 20,
        *,
        recent: bool = False,
        large: bool = False,
    ) -> Dict[str, object]:
            try:
                folder = self._resolve_folder_target(folder_text or "home")
            except ValueError as exc:
                return {"success": False, "action": "find_files", "message": str(exc)}

            if not folder.exists():
                return {"success": False, "action": "find_files", "message": f"{self._display_target_name(folder)} does not exist."}
            if not folder.is_dir():
                return {"success": False, "action": "find_files", "message": "That is a file, not a folder."}

            query = self._sanitize_file_reference(query_text).lower()
            query = "" if query == "*" else query
            extension = ""
            extension_aliases = {
                "pdf": ".pdf",
                "pdfs": ".pdf",
                "document": ".docx",
                "documents": ".docx",
                "word": ".docx",
                "text": ".txt",
                "notes": ".txt",
                "image": ".jpg",
                "images": ".jpg",
                "photo": ".jpg",
                "photos": ".jpg",
                "video": ".mp4",
                "videos": ".mp4",
            }
            for token, ext in extension_aliases.items():
                if re.search(rf"\b{re.escape(token)}\b", query):
                    extension = ext
                    break
            ext_match = re.search(r"\.([a-z0-9]{1,8})\b", query)
            if ext_match:
                extension = "." + ext_match.group(1)

            name_query = re.sub(
                r"\b(find|search|look|for|about|my|laptop|computer|pc|files?|all|the|named|called|documents?|pdfs?|images?|photos?|videos?|music|text|notes|recent|recently|modified|latest|newest|large|largest|biggest)\b",
                " ",
                query,
            )
            name_query = re.sub(r"\.[a-z0-9]{1,8}\b", " ", name_query)
            name_query = re.sub(r"\s+", " ", name_query).strip()

            results: list[dict[str, object]] = []
            scanned = 0
            partial = False
            started_at = time.monotonic()
            try:
                for item in self._iter_search_files(folder):
                    if time.monotonic() - started_at > MAX_FILE_SEARCH_SECONDS:
                        partial = True
                        break
                    scanned += 1
                    if scanned > MAX_FILE_SEARCH_SCANNED:
                        partial = True
                        break
                    suffix = item.suffix.lower()
                    if extension and suffix != extension:
                        continue
                    match_type = ""
                    lowered_name = item.name.lower()
                    if name_query and name_query in lowered_name:
                        match_type = "name"
                    elif not name_query:
                        match_type = "name"
                    elif self._content_matches(item, name_query):
                        match_type = "content"
                    else:
                        continue
                    try:
                        stat = item.stat()
                    except Exception:
                        continue
                    results.append(
                        {
                            "index": 0,
                            "name": item.name,
                            "path": str(item),
                            "parent": str(item.parent),
                            "size_bytes": int(stat.st_size),
                            "size": self._format_size(stat.st_size),
                            "modified_at": stat.st_mtime,
                            "match_type": match_type,
                        }
                    )
                    if not recent and not large and len(results) >= limit:
                        break
            except PermissionError:
                return {"success": False, "action": "find_files", "message": f"Permission denied while searching {folder}."}
            except Exception as exc:
                return {"success": False, "action": "find_files", "message": f"Could not search {folder}: {exc}"}

            if recent:
                results.sort(key=lambda item: float(item.get("modified_at") or 0.0), reverse=True)
            elif large:
                results.sort(key=lambda item: int(item.get("size_bytes") or 0), reverse=True)
            results = results[:limit]
            for index, item in enumerate(results, start=1):
                item["index"] = index

            self._last_folder_target = folder
            self._last_file_search_results = list(results)
            if len(results) == 1:
                self._last_file_target = Path(str(results[0]["path"]))
            if not results:
                suffix = " Search stopped early, so results may be incomplete." if partial else ""
                return {
                    "success": True,
                    "action": "find_files",
                    "message": f"No matching files found in {self._display_target_name(folder)}.{suffix}",
                    "data": {"results": [], "folder": str(folder), "query": query_text, "partial": partial, "scanned": scanned},
                    "partial": partial,
                }
            lines = [
                f"{item['index']}. {item['name']} ({item['size']}, {item['match_type']}) - {item['parent']}"
                for item in results
            ]
            suffix = "\nSearch stopped early, so these are partial results." if partial else ""
            return {
                "success": True,
                "action": "find_files",
                "message": "Found files:\n" + "\n".join(lines) + suffix,
                "data": {"results": results, "folder": str(folder), "query": query_text, "partial": partial, "scanned": scanned},
                "results": results,
                "partial": partial,
                "query": query_text,
            }


    def _iter_search_files(self, folder: Path) -> Iterable[Path]:
            base_dir = _runtime_base_dir().resolve()
            for root, dirnames, filenames in os.walk(folder):
                root_path = Path(root)
                if self._is_search_excluded_path(root_path, base_dir=base_dir):
                    dirnames[:] = []
                    continue
                dirnames[:] = [
                    name
                    for name in dirnames
                    if not self._is_search_excluded_path(root_path / name, base_dir=base_dir)
                ]
                for filename in filenames:
                    path = root_path / filename
                    if not self._is_search_excluded_path(path, base_dir=base_dir):
                        yield path


    def _is_search_excluded_path(self, path: Path, *, base_dir: Path) -> bool:
            try:
                resolved = path.resolve(strict=False)
            except Exception:
                resolved = path
            try:
                if self._is_protected_path(resolved):
                    return True
            except Exception:
                return True
            parts = {part.lower() for part in resolved.parts}
            if parts & SEARCH_EXCLUDED_DIR_NAMES:
                return True
            if any(part.startswith("pytest-cache-files") for part in parts):
                return True
            runtime_exclusions = (
                base_dir / "database",
                base_dir / ".continue",
                base_dir / ".git",
                base_dir / ".pytest_cache",
                base_dir / "node_modules",
                base_dir / "venv",
                base_dir / ".venv",
                base_dir / "tests" / "_tmp",
            )
            return any(self._is_relative_to(resolved, excluded) for excluded in runtime_exclusions)


    def _content_matches(self, path: Path, query: str) -> bool:
            if not query or path.suffix.lower() not in SEARCH_TEXT_EXTENSIONS:
                return False
            try:
                if path.stat().st_size > MAX_CONTENT_SEARCH_BYTES:
                    return False
                return query in path.read_text(encoding="utf-8", errors="ignore").lower()
            except Exception:
                return False


    @staticmethod
    def _is_relative_to(path: Path, root: Path) -> bool:
            try:
                path.resolve(strict=False).relative_to(root.resolve(strict=False))
                return True
            except Exception:
                return False


    def _read_supported_file_content(self, path: Path, *, max_chars: int) -> Dict[str, object]:
            try:
                if path.stat().st_size > 2_000_000:
                    return {"success": False, "action": "read_file", "message": "That file is too large to read safely in chat."}
            except Exception as exc:
                return {"success": False, "action": "read_file", "message": f"Could not inspect {self._display_file_name(path)}: {exc}"}

            suffix = path.suffix.lower()
            try:
                if suffix == ".csv":
                    content = self._preview_csv(path)
                    reader = "csv"
                elif suffix == ".pdf":
                    content = self._read_pdf(path)
                    reader = "pdf"
                elif suffix == ".docx":
                    content = self._read_docx(path)
                    reader = "docx"
                elif suffix == ".xlsx":
                    content = self._preview_xlsx(path)
                    reader = "xlsx"
                elif suffix in READ_TEXT_EXTENSIONS or not suffix:
                    content = path.read_text(encoding="utf-8", errors="replace")
                    reader = "text"
                else:
                    return {
                        "success": False,
                        "action": "read_file",
                        "message": f"{self._display_file_name(path)} has unsupported file type {suffix or 'unknown'}.",
                        "status": "unsupported_file_type",
                    }
            except ImportError as exc:
                return {
                    "success": False,
                    "action": "read_file",
                    "message": str(exc),
                    "status": "setup_needed",
                }
            except Exception as exc:
                return {"success": False, "action": "read_file", "message": f"Could not read {self._display_file_name(path)}: {exc}"}

            if len(content) > max_chars:
                content = content[:max_chars].rstrip() + f"\n\n... truncated at {max_chars} characters."
            return {"success": True, "content": content, "reader": reader}


    def _preview_csv(self, path: Path, *, rows: int = 12) -> str:
            with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
                reader = csv.reader(handle)
                lines = []
                for index, row in enumerate(reader):
                    if index >= rows:
                        break
                    lines.append(", ".join(row))
            return "\n".join(lines)


    def _read_pdf(self, path: Path) -> str:
            try:
                from pypdf import PdfReader  # type: ignore
            except Exception:
                try:
                    from PyPDF2 import PdfReader  # type: ignore
                except Exception as exc:
                    raise ImportError("PDF reading needs pypdf or PyPDF2 installed.") from exc
            reader = PdfReader(str(path))
            pages = []
            for page in list(reader.pages)[:5]:
                pages.append(str(page.extract_text() or "").strip())
            return "\n\n".join(page for page in pages if page)


    def _read_docx(self, path: Path) -> str:
            try:
                import docx  # type: ignore
            except Exception as exc:
                raise ImportError("DOCX reading needs python-docx installed.") from exc
            document = docx.Document(str(path))
            return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)


    def _preview_xlsx(self, path: Path, *, rows: int = 12, columns: int = 8) -> str:
            try:
                import openpyxl  # type: ignore
            except Exception as exc:
                raise ImportError("XLSX reading needs openpyxl installed.") from exc
            workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            sheet = workbook.active
            lines = []
            for row in sheet.iter_rows(max_row=rows, max_col=columns, values_only=True):
                lines.append(", ".join("" if value is None else str(value) for value in row))
            workbook.close()
            return "\n".join(lines)


    def _largest_files(self, folder_text: str = "home", limit: int = 10) -> Dict[str, str | bool]:
            try:
                folder = self._resolve_folder_target(folder_text or "home")
            except ValueError as exc:
                return {"success": False, "action": "largest_files", "message": str(exc)}

            if not folder.exists():
                return {"success": False, "action": "largest_files", "message": f"{self._display_target_name(folder)} does not exist."}
            if not folder.is_dir():
                return {"success": False, "action": "largest_files", "message": "That is a file, not a folder."}

            candidates: list[tuple[int, Path]] = []
            scanned = 0
            try:
                for item in folder.rglob("*"):
                    scanned += 1
                    if scanned > 8000:
                        break
                    if item.is_file():
                        try:
                            candidates.append((item.stat().st_size, item))
                        except Exception:
                            continue
            except PermissionError:
                return {"success": False, "action": "largest_files", "message": f"Permission denied while scanning {folder}."}
            except Exception as exc:
                return {"success": False, "action": "largest_files", "message": f"Could not scan {folder}: {exc}"}

            candidates.sort(reverse=True, key=lambda pair: pair[0])
            lines = [f"{path.name} ({self._format_size(size)}) - {path.parent}" for size, path in candidates[:limit]]
            if not lines:
                return {"success": True, "action": "largest_files", "message": f"No files found in {self._display_target_name(folder)}."}
            return {"success": True, "action": "largest_files", "message": "Largest files:\n" + "\n".join(lines)}


    def _organize_folder_preview(self, folder_text: str = "downloads") -> Dict[str, str | bool]:
            try:
                folder = self._resolve_folder_target(folder_text or "downloads")
            except ValueError as exc:
                return {"success": False, "action": "organize_folder_preview", "message": str(exc)}

            if not folder.exists():
                return {"success": False, "action": "organize_folder_preview", "message": f"{self._display_target_name(folder)} does not exist."}
            if not folder.is_dir():
                return {"success": False, "action": "organize_folder_preview", "message": "That is a file, not a folder."}

            categories = {
                "Images": {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".svg"},
                "Documents": {".pdf", ".doc", ".docx", ".txt", ".md", ".rtf", ".csv", ".xlsx", ".pptx"},
                "Archives": {".zip", ".rar", ".7z", ".tar", ".gz"},
                "Installers": {".exe", ".msi", ".apk", ".dmg", ".pkg"},
                "Videos": {".mp4", ".mov", ".avi", ".mkv", ".webm"},
                "Audio": {".mp3", ".wav", ".flac", ".m4a", ".ogg"},
            }
            counts = {name: 0 for name in categories}
            counts["Other"] = 0

            try:
                for item in folder.iterdir():
                    if not item.is_file():
                        continue
                    suffix = item.suffix.lower()
                    matched = False
                    for category, suffixes in categories.items():
                        if suffix in suffixes:
                            counts[category] += 1
                            matched = True
                            break
                    if not matched:
                        counts["Other"] += 1
            except Exception as exc:
                return {"success": False, "action": "organize_folder_preview", "message": f"Could not inspect {folder}: {exc}"}

            lines = [f"{name}: {count}" for name, count in counts.items() if count]
            if not lines:
                message = f"{self._display_target_name(folder)} has no files to organize."
            else:
                message = (
                    f"Organization preview for {self._display_target_name(folder)}:\n"
                    + "\n".join(lines)
                    + "\nNo files were moved."
                )
            self._last_folder_target = folder
            return {"success": True, "action": "organize_folder_preview", "message": message}


    def _update_file(self, path_text: str, content: str, append: bool = True) -> Dict[str, str | bool]:
            if not content:
                return {
                    "success": False,
                    "action": "update_file",
                    "message": "Tell me what text you want me to add.",
                }

            try:
                path = self._resolve_file_target(path_text)
            except ValueError as exc:
                return {"success": False, "action": "update_file", "message": str(exc)}

            if path.exists() and path.is_dir():
                return {
                    "success": False,
                    "action": "update_file",
                    "message": "That is a folder, not a file.",
                }

            path.parent.mkdir(parents=True, exist_ok=True)

            existing = ""
            if path.exists():
                try:
                    existing = path.read_text(encoding="utf-8")
                except Exception:
                    existing = ""

            if append and existing:
                separator = "" if existing.endswith(("\n", "\r")) else "\n"
                new_content = existing + separator + content
            elif append and not existing:
                new_content = content
            else:
                new_content = content

            try:
                path.write_text(new_content, encoding="utf-8")
            except Exception as exc:
                return {
                    "success": False,
                    "action": "update_file",
                    "message": f"Could not update {self._display_file_name(path)} at {path}: {exc}",
                }

            self._last_file_target = path

            return {
                "success": True,
                "action": "update_file",
                "message": f"Updated {self._display_file_name(path)} in {self._display_parent_name(path)}.",
            }


    def _show_last_file_path(self) -> Dict[str, str | bool]:
            if self._last_file_target is None:
                return {
                    "success": False,
                    "action": "show_path",
                    "message": "I don't know which file you mean yet.",
                }
            return {
                "success": True,
                "action": "show_path",
                "message": f"{self._display_file_name(self._last_file_target)} is at {self._last_file_target}.",
            }


    def _delete_file(self, path_text: str) -> Dict[str, str | bool]:
            try:
                path = self._resolve_existing_target(path_text, target_kind="file")
            except ValueError as exc:
                return {"success": False, "action": "delete_file", "message": str(exc)}

            if not path.exists():
                return {
                    "success": False,
                    "action": "delete_file",
                    "message": f"{self._display_file_name(path)} does not exist.",
                }

            if path.is_dir():
                return {
                    "success": False,
                    "action": "delete_file",
                    "message": "That is a folder, not a file.",
                }

            self._pending_delete_target = path
            return {
                "success": False,
                "action": "delete_file",
                "message": f"Do you want me to delete {self._display_file_name(path)}? Reply yes or no.",
            }


    def _delete_folder(self, path_text: str) -> Dict[str, str | bool]:
            try:
                path = self._resolve_existing_target(path_text, target_kind="folder")
            except ValueError as exc:
                return {"success": False, "action": "delete_folder", "message": str(exc)}

            if not path.exists():
                return {
                    "success": False,
                    "action": "delete_folder",
                    "message": f"{self._display_target_name(path)} does not exist.",
                }

            if not path.is_dir():
                return {
                    "success": False,
                    "action": "delete_folder",
                    "message": "That is a file, not a folder.",
                }

            self._pending_delete_target = path
            self._last_folder_target = path
            return {
                "success": False,
                "action": "delete_folder",
                "message": f"Do you want me to delete folder {self._display_target_name(path)}? Reply yes or no.",
            }


    def _rename_target(self, source_text: str, new_name: str, target_kind: str = "") -> Dict[str, str | bool]:
            try:
                source = self._resolve_existing_target(source_text, target_kind=target_kind)
            except ValueError as exc:
                return {"success": False, "action": "rename", "message": str(exc)}

            if not source.exists():
                return {
                    "success": False,
                    "action": "rename",
                    "message": f"{self._display_target_name(source)} does not exist.",
                }

            clean_name = self._clean_file_name(new_name)
            if not clean_name:
                return {
                    "success": False,
                    "action": "rename",
                    "message": "Tell me the new name you want to use.",
                }

            if "\\" in clean_name or "/" in clean_name:
                return {
                    "success": False,
                    "action": "rename",
                    "message": "For renaming, give me just the new name, not a full path.",
                }

            destination = source.with_name(clean_name)
            if self._is_protected_path(destination):
                return {
                    "success": False,
                    "action": "rename",
                    "message": "That location is protected. Jarvis cannot rename items inside Windows or critical system folders.",
                }

            if destination.exists():
                return {
                    "success": False,
                    "action": "rename",
                    "message": f"{self._display_target_name(destination)} already exists at {destination}.",
                }

            try:
                source.rename(destination)
            except Exception as exc:
                return {
                    "success": False,
                    "action": "rename",
                    "message": f"Could not rename {self._display_target_name(source)}: {exc}",
                }

            self._remember_target(destination)
            return {
                "success": True,
                "action": "rename",
                "message": f"Renamed {self._display_target_name(source)} to {self._display_target_name(destination)}.",
            }


    def _move_target(self, source_text: str, destination_text: str, target_kind: str = "") -> Dict[str, str | bool]:
            try:
                source = self._resolve_existing_target(source_text, target_kind=target_kind)
            except ValueError as exc:
                return {"success": False, "action": "move", "message": str(exc)}

            if not source.exists():
                return {
                    "success": False,
                    "action": "move",
                    "message": f"{self._display_target_name(source)} does not exist.",
                }

            try:
                destination = self._resolve_move_destination(source, destination_text)
            except ValueError as exc:
                return {"success": False, "action": "move", "message": str(exc)}

            if destination.exists():
                return {
                    "success": False,
                    "action": "move",
                    "message": f"{self._display_target_name(destination)} already exists at {destination}.",
                }

            try:
                self.local_files_connector.move(source, destination)
            except Exception as exc:
                return {
                    "success": False,
                    "action": "move",
                    "message": f"Could not move {self._display_target_name(source)}: {exc}",
                }

            moved_path = Path(destination)
            self._remember_target(moved_path)
            return {
                "success": True,
                "action": "move",
                "message": f"Moved {self._display_target_name(moved_path)} to {moved_path}.",
            }


    def _resolve_move_destination(self, source: Path, destination_text: str) -> Path:
            cleaned = self._sanitize_file_reference(destination_text)
            if not cleaned:
                raise ValueError("Tell me where you want me to move it.")

            folder_target = self._resolve_folder_phrase(cleaned)
            if folder_target is not None:
                if self._is_protected_path(folder_target):
                    raise ValueError("That location is protected. Jarvis cannot access Windows or critical system folders.")
                return folder_target / source.name

            destination = self._resolve_laptop_path(cleaned)
            if destination.exists() and destination.is_dir():
                return destination / source.name

            if cleaned.endswith(("\\", "/")):
                return destination / source.name

            return destination


    def _resolve_folder_phrase(self, path_text: str) -> Path | None:
            raw_cleaned = self._sanitize_file_reference(path_text)
            explicit_folder_match = re.match(
                r"^(?:on|in|at|inside|under)\s+(?:the\s+)?(?:folder|directory)\s+(.+)$|^(?:the\s+)?(?:folder|directory)\s+(.+)$",
                raw_cleaned,
                flags=re.IGNORECASE,
            )
            if explicit_folder_match:
                explicit_value = (explicit_folder_match.group(1) or explicit_folder_match.group(2) or "").strip()
                if explicit_value:
                    return self._resolve_path(explicit_value).resolve()

            cleaned = self._clean_location_phrase(path_text)
            if not cleaned:
                return None

            alias_path = self._resolve_user_alias_path(cleaned)
            if alias_path is not None:
                lowered = cleaned.lower()
                if lowered in self.USER_PATH_ALIASES or lowered.startswith("the "):
                    return alias_path

            lowered = cleaned.lower()
            if lowered in self.USER_PATH_ALIASES:
                return self.USER_PATH_ALIASES[lowered]

            if self._looks_like_windows_absolute_path(cleaned) and cleaned.endswith(("\\", "/")):
                return Path(cleaned.replace("/", "\\"))

            return None


    def _handle_delete_confirmation(self, command: str) -> Dict[str, str | bool]:
            target = self._pending_delete_target
            response = command.strip().lower().rstrip(".!?")

            if response in {"yes", "y", "delete it", "confirm", "go ahead"}:
                self._pending_delete_target = None
                if not target or not target.exists():
                    return {
                        "success": False,
                        "action": "delete",
                        "message": "That item is no longer available to delete.",
                    }
                action = "delete_folder" if target.is_dir() else "delete_file"
                plan_command = f"delete {'folder' if target.is_dir() else 'file'} {target}"
                executor = ToolExecutor(registry=self._build_automation_tool_registry(), enforce_policy=True)
                return executor.execute(
                    ActionPlan(
                        original_text=plan_command,
                        steps=[
                            ActionStep(
                                step_id="step1",
                                tool_name="file",
                                intent="file",
                                action=action,
                                args={"path": str(target), "confirmed": True},
                            )
                        ],
                        is_multistep=False,
                    ),
                    ToolContext(
                        command=plan_command,
                        intent="file",
                        session_id=self._active_session_id,
                        request_id=self._active_turn_id,
                        payload={"turn_id": self._active_turn_id} if self._active_turn_id else {},
                        confirmation_state={"confirmed": True},
                        security_state={"step_up_verified": self._active_step_up_verified},
                    ),
                )

            if response in {"no", "n", "cancel", "stop", "don't", "do not"}:
                self._pending_delete_target = None
                return {
                    "success": False,
                    "action": "delete",
                    "message": "Deletion cancelled.",
                }

            return {
                "success": False,
                "action": "delete",
                "message": f"Please reply yes or no to confirm deleting {self._display_target_name(target)}.",
            }


    def _is_protected_path(self, path: Path) -> bool:
            try:
                resolved = path.resolve()
            except Exception:
                resolved = path

            resolved_text = str(resolved).lower()
            drive_root_match = re.match(r"^[a-z]:\\?$", resolved_text)
            if drive_root_match:
                return True

            workspace_text = str(_runtime_base_dir().resolve()).lower()
            if resolved_text.startswith(workspace_text):
                return False

            for prefix in self.PROTECTED_PATH_PREFIXES:
                prefix_text = str(prefix.resolve()).lower()
                if resolved_text == prefix_text or resolved_text.startswith(prefix_text + "\\"):
                    return True

            for pattern in self.PROTECTED_PATH_PATTERNS:
                if pattern.match(resolved_text):
                    return True

            return False


